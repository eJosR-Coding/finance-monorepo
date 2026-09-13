"""Build the Excel workbook that cross-checks the engine.

The point is NOT to paste our numbers into a spreadsheet - that would prove
nothing. Every cell here is a live Excel formula, so the schedule recomputes
inside Excel and a comparison block subtracts it from what the system produced.
If the difference column is not zero, one of the two is wrong.

Formula names are written in English (PMT, ROUND). Excel stores them that way
and renders them localized, so a Spanish Excel shows PAGO and REDONDEAR.

Run it with:  uv run python scripts/generate_excel.py
"""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

from app.core.config import BASE_DIR, settings  # noqa: E402
from app.services.finance import Schedule, build_schedule  # noqa: E402
from scripts.academic_cases import ACADEMIC_CASES, AcademicCase  # noqa: E402

OUTPUT = BASE_DIR.parent.parent / "docs" / "evidence" / "prestameami-cronogramas.xlsx"

INK = "FF201E1D"
ACCENT = "FFEC3013"
SURFACE = "FFEAE9E9"
HEADING = Font(name="Calibri", bold=True, color="FFFFFFFF", size=11)
LABEL = Font(name="Calibri", bold=True, size=10)
TITLE = Font(name="Calibri", bold=True, size=14, color=INK)
MONEY = '"S/" #,##0.00'
RATE9 = "0.000000000"
PERCENT = "0.00%"
THIN = Side(style="thin", color="FFBAB6B6")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _header(sheet, row: int, labels: list[str]) -> None:
    for index, text in enumerate(labels, start=1):
        cell = sheet.cell(row=row, column=index, value=text)
        cell.font = HEADING
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BOX


def _label(sheet, row: int, text: str, value=None, number_format: str | None = None):
    sheet.cell(row=row, column=1, value=text).font = LABEL
    cell = sheet.cell(row=row, column=2, value=value)
    if number_format:
        cell.number_format = number_format
    cell.fill = PatternFill("solid", fgColor=SURFACE)
    cell.border = BOX
    return cell


def _schedule_block(sheet, first_row: int, count: int, dynamic: bool) -> None:
    """Write the amortization rows as formulas.

    `dynamic` builds rows that hide themselves when the installment number goes
    past n, so the free simulator sheet can have a fixed 14 rows.
    """
    for offset in range(count):
        row = first_row + offset
        number = offset + 1
        previous = row - 1
        guard = f"IF(A{row}>$B$7,\"\","
        close = ")" if dynamic else ""
        prefix = guard if dynamic else ""

        sheet.cell(row=row, column=1, value=number).border = BOX

        date_cell = sheet.cell(row=row, column=2, value=f"={prefix}$B$10+A{row}*$B$8{close}")
        date_cell.number_format = "DD/MM/YYYY"

        opening = "$B$4" if number == 1 else f"G{previous}"
        sheet.cell(row=row, column=3, value=f"={prefix}{opening}{close}")

        sheet.cell(row=row, column=4, value=f"={prefix}ROUND(C{row}*$B$12,2){close}")

        if dynamic:
            # The last installment absorbs the rounding residue.
            sheet.cell(
                row=row,
                column=5,
                value=f"={guard}IF(A{row}=$B$7,C{row},ROUND($B$13-D{row},2)))",
            )
            sheet.cell(
                row=row,
                column=6,
                value=f"={guard}IF(A{row}=$B$7,ROUND(D{row}+E{row},2),$B$13))",
            )
        elif number == count:
            sheet.cell(row=row, column=5, value=f"=C{row}")
            sheet.cell(row=row, column=6, value=f"=ROUND(D{row}+E{row},2)")
        else:
            sheet.cell(row=row, column=5, value=f"=ROUND($B$13-D{row},2)")
            sheet.cell(row=row, column=6, value="=$B$13")

        sheet.cell(row=row, column=7, value=f"={prefix}ROUND(C{row}-E{row},2){close}")

        for column in range(1, 8):
            cell = sheet.cell(row=row, column=column)
            cell.border = BOX
            if column >= 3:
                cell.number_format = MONEY


def _inputs(sheet, case: AcademicCase | None, count: int) -> None:
    sheet["A3"] = "DATOS DE ENTRADA"
    sheet["A3"].font = Font(bold=True, color=ACCENT, size=10)

    _label(sheet, 4, "Capital (P)", float(case.amount) if case else 140.00, MONEY)
    _label(sheet, 5, "TEA", float(case.annual_rate) / 100 if case else 0.40, PERCENT)
    _label(sheet, 6, "Plazo (dias)", case.term_days if case else 14)
    _label(sheet, 7, "Numero de cuotas (n)", case.installments_count if case else count)
    _label(sheet, 8, "Frecuencia (dias)", case.payment_frequency_days if case else 1)
    _label(sheet, 9, "Base anual (dias)", settings.days_per_year)
    date_cell = _label(sheet, 10, "Fecha de inicio", case.start_date if case else None)
    date_cell.number_format = "DD/MM/YYYY"

    sheet["A11"] = "CALCULOS"
    sheet["A11"].font = Font(bold=True, color=ACCENT, size=10)

    rate = _label(sheet, 12, "Tasa periodica (i)", "=ROUND((1+B5)^(B8/B9)-1,9)", RATE9)
    rate.font = Font(bold=True)
    _label(sheet, 13, "Cuota (C)", "=ROUND(PMT(B12,B7,-B4),2)", MONEY)

    last = 18 + count - 1
    _label(sheet, 14, "Interes total", f"=ROUND(SUM(D18:D{last}),2)", MONEY)
    _label(sheet, 15, "Total a pagar", f"=ROUND(SUM(F18:F{last}),2)", MONEY)

    sheet["D4"] = "i = (1 + TEA)^(dias / base) - 1"
    sheet["D5"] = "C = PAGO(i; n; -P)     [PMT en Excel en ingles]"
    sheet["D6"] = "Interes_t = REDONDEAR(saldo_anterior * i; 2)"
    sheet["D7"] = "Amortizacion_t = REDONDEAR(C - Interes_t; 2)"
    sheet["D8"] = "Saldo_t = REDONDEAR(saldo_anterior - Amortizacion_t; 2)"
    sheet["D9"] = "Ultima cuota: absorbe el residuo para cerrar en 0.00"
    for row in range(4, 10):
        sheet.cell(row=row, column=4).font = Font(italic=True, size=9, color="FF605D5D")


def _comparison(sheet, schedule: Schedule, first_row: int, count: int) -> None:
    """System value vs Excel formula vs live difference."""
    start = first_row + count + 2
    sheet.cell(row=start, column=1, value="COMPARACION SISTEMA vs EXCEL").font = Font(
        bold=True, color=ACCENT, size=10
    )
    _header(sheet, start + 1, ["Variable", "Sistema", "Excel", "Diferencia", "Resultado"])

    last = first_row + count - 1
    rows = [
        ("Tasa periodica (i)", float(schedule.periodic_rate), "=B12", RATE9),
        ("Cuota (C)", float(schedule.installment_amount), "=B13", MONEY),
        ("Interes total", float(schedule.total_interest), "=B14", MONEY),
        ("Total a pagar", float(schedule.total_payment), "=B15", MONEY),
        ("Saldo final", 0.00, f"=G{last}", MONEY),
    ]
    for offset, (name, system, formula, number_format) in enumerate(rows):
        row = start + 2 + offset
        sheet.cell(row=row, column=1, value=name).font = LABEL
        sheet.cell(row=row, column=2, value=system).number_format = number_format
        sheet.cell(row=row, column=3, value=formula).number_format = number_format
        sheet.cell(row=row, column=4, value=f"=C{row}-B{row}").number_format = number_format
        sheet.cell(row=row, column=5, value=f'=IF(ROUND(D{row},9)=0,"OK","REVISAR")')
        for column in range(1, 6):
            sheet.cell(row=row, column=column).border = BOX

    note = start + 2 + len(rows) + 1
    sheet.cell(
        row=note,
        column=1,
        value=(
            "La columna Excel son formulas vivas: cambia un dato de entrada y todo "
            "se recalcula. La columna Sistema es lo que devuelve la API."
        ),
    ).font = Font(italic=True, size=9, color="FF605D5D")


def _widths(sheet) -> None:
    for column, width in zip("ABCDEFG", [26, 16, 16, 14, 16, 14, 16], strict=False):
        sheet.column_dimensions[column].width = width


def _case_sheet(workbook: Workbook, case: AcademicCase) -> None:
    schedule = build_schedule(
        amount=case.amount,
        rate_type=case.rate_type,
        annual_rate=case.annual_rate,
        start_date=case.start_date,
        term_days=case.term_days,
        installments_count=case.installments_count,
        payment_frequency_days=case.payment_frequency_days,
    )
    count = len(schedule.rows)
    sheet = workbook.create_sheet(case.key.replace("caso-", "Caso "))

    sheet["A1"] = case.title
    sheet["A1"].font = TITLE
    _inputs(sheet, case, count)
    _header(
        sheet,
        17,
        ["#", "Fecha", "Saldo inicial", "Interes", "Amortizacion", "Cuota", "Saldo final"],
    )
    _schedule_block(sheet, 18, count, dynamic=False)
    _comparison(sheet, schedule, 18, count)
    _widths(sheet)


def _simulator_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Simulador libre")
    sheet["A1"] = "Simulador libre — cambia los datos de entrada y recalcula"
    sheet["A1"].font = TITLE
    _inputs(sheet, None, 14)
    sheet["B4"] = 140.00
    sheet["B7"] = 2
    sheet["B8"] = 7
    sheet["B10"] = ACADEMIC_CASES[0].start_date
    sheet["B10"].number_format = "DD/MM/YYYY"
    _header(
        sheet,
        17,
        ["#", "Fecha", "Saldo inicial", "Interes", "Amortizacion", "Cuota", "Saldo final"],
    )
    _schedule_block(sheet, 18, 14, dynamic=True)
    sheet.cell(
        row=34,
        column=1,
        value=(
            "Limites del producto: monto maximo S/ 200.00, plazo maximo 14 dias. "
            "Las filas sobrantes se ocultan solas cuando n es menor a 14."
        ),
    ).font = Font(italic=True, size=9, color="FF605D5D")
    _widths(sheet)


def _readme_sheet(workbook: Workbook) -> None:
    sheet = workbook.active
    sheet.title = "Instrucciones"
    lines = [
        ("Prestameami.pe — verificacion del cronograma en Excel", TITLE),
        ("", None),
        ("Para que sirve", Font(bold=True, color=ACCENT)),
        ("Comprobar, fuera del sistema, que el cronograma que calcula la aplicacion", None),
        ("es correcto. Ninguna celda tiene numeros pegados a mano: todas son formulas", None),
        ("de Excel, asi que el cronograma se recalcula dentro de Excel.", None),
        ("", None),
        ("Como leerlo", Font(bold=True, color=ACCENT)),
        ("1. Las hojas 'Caso 1' y 'Caso 2' son los casos del Anexo C.", None),
        ("2. Abajo de cada cronograma hay una tabla que compara:", None),
        ("      Sistema    = lo que devuelve la API", None),
        ("      Excel      = lo que calcula la formula de esta hoja", None),
        ("      Diferencia = Excel menos Sistema, debe dar 0.00", None),
        ("3. 'Simulador libre' deja cambiar capital, TEA, cuotas y frecuencia.", None),
        ("", None),
        ("Formulas usadas", Font(bold=True, color=ACCENT)),
        ("   i = (1 + TEA) ^ (dias del periodo / 360) - 1          (9 decimales)", None),
        ("   C = PAGO(i; n; -P)                                    (2 decimales)", None),
        ("   Interes_t      = REDONDEAR(saldo anterior * i; 2)", None),
        ("   Amortizacion_t = REDONDEAR(C - Interes_t; 2)", None),
        ("   Saldo_t        = REDONDEAR(saldo anterior - Amortizacion_t; 2)", None),
        ("", None),
        ("Por que la ultima cuota es distinta", Font(bold=True, color=ACCENT)),
        ("Al redondear cada periodo a 2 decimales queda un residuo de centimos.", None),
        ("La ultima cuota amortiza todo el saldo que queda, asi el cronograma cierra", None),
        ("en S/ 0.00 exacto. Por eso en el Caso 2 la ultima cuota es S/ 60.60 y no", None),
        ("S/ 60.59: ese centimo es el ajuste, no un error.", None),
        ("", None),
        ("Nota sobre el idioma de Excel", Font(bold=True, color=ACCENT)),
        ("Las formulas se guardan en ingles (PMT, ROUND) y Excel las muestra en el", None),
        ("idioma que tengas instalado (PAGO, REDONDEAR). Es normal.", None),
    ]
    for index, (text, font) in enumerate(lines, start=1):
        cell = sheet.cell(row=index, column=1, value=text)
        if font is not None:
            cell.font = font
    sheet.column_dimensions["A"].width = 95


def main() -> None:
    workbook = Workbook()
    _readme_sheet(workbook)
    for case in ACADEMIC_CASES:
        _case_sheet(workbook, case)
    _simulator_sheet(workbook)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUTPUT)
    print(f"Generado: {OUTPUT}")
    print(f"Hojas: {', '.join(workbook.sheetnames)}")


if __name__ == "__main__":
    main()
